"""Build the Phase 2 Word deliverables from the committed Markdown sources.

    python report/build_docx.py

Produces, next to the sources:
    report/phase2_report.docx
    report/annotated_bibliography.docx

Layout is fixed here so the page budget is reproducible: US Letter, 0.75in margins,
Calibri 10pt body, 8.5pt tables, Consolas 8.5pt code, body figures 4.4in wide.
Display equations are rasterised with matplotlib mathtext (python-docx has no OMML
support); inline maths is mapped to Unicode through MATH explicitly, so an unmapped
span raises rather than silently emitting LaTeX source.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
EQ_DIR = REPO / "figures" / "_eq"

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
BODY_PT = 10.0
TABLE_PT = 8.5
CODE_PT = 8.5
FIG_W_BODY = 4.4
FIG_W_APPX = 5.3
TEXT_W = 7.0  # 8.5in letter - 2 * 0.75in margin

CURRENCY = "\x00CUR\x00"  # protects "NT$" from the maths scanner

# Every inline maths span that appears in the sources, mapped explicitly.
MATH = {
    "K": "K",
    "t": "t",
    "H": "H",
    "n": "n",
    "8": "8",
    "-2": "\u22122",
    "(0,1)": "(0,1)",
    "(G, H)": "(G, H)",
    "O(n)": "O(n)",
    r"O(n \log n)": "O(n log n)",
    "1 - 0.2211": "1 \u2212 0.2211",
    "3^8 = 6{,}561": "3\u2078 = 6,561",
    "n = 29{,}944": "n = 29,944",
    r"\lambda": "\u03bb",
    r"\gamma": "\u03b3",
    r"\alpha": "\u03b1",
    r"\eta": "\u03b7",
    r"\gamma T": "\u03b3T",
    r"\lambda = 34.3": "\u03bb = 34.3",
    r"\pm\infty": "\u00b1\u221e",
    r"\sqrt{n}": "\u221an",
    r"\sum h_i": "\u03a3h\u1d62",
    "p^* = 1/(1+r)": "p* = 1/(1+r)",
    "h_i = p_i(1-p_i)": "h\u1d62 = p\u1d62(1\u2212p\u1d62)",
    r"g_i = \partial_{\hat{y}} l": "g\u1d62 = \u2202l/\u2202\u0177",
    r"h_i = \partial^2_{\hat{y}} l": "h\u1d62 = \u2202\u00b2l/\u2202\u0177\u00b2",
    r"\hat{y}_i = \sigma(\sum_k f_k(\mathbf{x}_i))": "\u0177\u1d62 = \u03c3(\u03a3\u2096 f\u2096(x\u1d62))",
    r"\Omega(f) = \gamma T + \tfrac{1}{2}\lambda \|w\|^2 + \alpha\|w\|_1":
        "\u03a9(f) = \u03b3T + \u00bd\u03bb\u2016w\u2016\u00b2 + \u03b1\u2016w\u2016\u2081",
}


# --------------------------------------------------------------------------- maths


def _mathtext(src: str) -> str:
    """Make a LaTeX display equation acceptable to matplotlib's mathtext parser."""
    s = src.strip()
    s = s.replace(r"\tfrac", r"\frac").replace(r"\dfrac", r"\frac")
    s = s.replace(r"\text{", r"\mathrm{")
    s = s.replace(r"\!", "")
    s = s.replace(r"\,", r"\ ")
    return s


def render_equation(src: str, idx: int) -> Path:
    """Rasterise one display equation at high resolution; return the PNG path."""
    EQ_DIR.mkdir(parents=True, exist_ok=True)
    out = EQ_DIR / f"eq{idx}.png"
    expr = f"${_mathtext(src)}$"
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, expr, fontsize=13, color="#1a1a1a")
    try:
        fig.savefig(out, dpi=400, transparent=True, bbox_inches="tight", pad_inches=0.04)
    except Exception as exc:  # pragma: no cover - surfaces a bad expression loudly
        plt.close(fig)
        raise SystemExit(f"mathtext failed on equation {idx}:\n  {expr}\n  {exc}")
    plt.close(fig)
    return out


# ------------------------------------------------------------------- docx plumbing


def shade(el, hexcolor: str) -> None:
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    el.append(sh)


def tight_cell_margins(table, px: int = 60) -> None:
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(px if side in ("left", "right") else 20))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tblPr.append(mar)


def repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def setup(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, attr, Inches(0.75))
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0
    pf.widow_control = True


# --------------------------------------------------------------------- inline runs

# Code and maths spans are lifted out before emphasis is scanned, for two reasons.
# They may legitimately contain `*` -- `results/*.json`, `$p^* = 1/(1+r)$` -- which
# would otherwise terminate an italic span early and swallow the markers. And they
# may sit *inside* bold or italic, which needs the emphasis walk to recurse rather
# than emit one flat run. Stashing makes both cases fall out of the same pass.
# Backslash escapes are stashed for the same reason: `*p*\* = 1/(1+r)` must not let
# the escaped asterisk open a second italic span and run away to the next one.
ATOM = re.compile(r"`[^`]+`|\$[^$\n]+?\$|\\.")
SLOT = re.compile(r"\x01(\d+)\x01")

EMPH = re.compile(
    r"(\*\*\*.+?\*\*\*)"              # bold italic
    r"|(\*\*.+?\*\*)"                 # bold
    r"|(\*[^*]+?\*)",                 # italic
    re.S,
)


def add_runs(par, text: str, *, bold=False, italic=False, size=None, mono=False):
    """Render markdown inline formatting into runs on `par`."""
    atoms: list[tuple[str, str]] = []

    def stash(m):
        tok = m.group(0)
        if tok.startswith("\\"):
            atoms.append(("lit", tok[1:]))
        else:
            atoms.append(("code" if tok.startswith("`") else "math", tok[1:-1]))
        return f"\x01{len(atoms) - 1}\x01"

    _emph(par, ATOM.sub(stash, text), atoms, bold, italic, size, mono)
    return par


def _emph(par, text, atoms, bold, italic, size, mono):
    """Walk `***`/`**`/`*` spans, recursing so nesting composes."""
    pos = 0
    for m in EMPH.finditer(text):
        if m.start() > pos:
            _atoms(par, text[pos:m.start()], atoms, bold, italic, size, mono)
        tok = m.group(0)
        if tok.startswith("***"):
            _emph(par, tok[3:-3], atoms, True, True, size, mono)
        elif tok.startswith("**"):
            _emph(par, tok[2:-2], atoms, True, italic, size, mono)
        else:
            _emph(par, tok[1:-1], atoms, bold, True, size, mono)
        pos = m.end()
    if pos < len(text):
        _atoms(par, text[pos:], atoms, bold, italic, size, mono)


def _atoms(par, text, atoms, bold, italic, size, mono):
    """Expand stashed code/maths slots, inheriting the surrounding emphasis."""
    pos = 0
    for m in SLOT.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], bold, italic, size, mono)
        kind, body = atoms[int(m.group(1))]
        if kind == "lit":
            _run(par, body, bold, italic, size, mono)
        elif kind == "code":
            r = _run(par, body, bold, italic, size, True)
            r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        else:
            if body not in MATH:
                raise SystemExit(f"unmapped inline maths span: {body!r}")
            _run(par, MATH[body], bold, True, size, mono)
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], bold, italic, size, mono)


def _run(par, txt, bold, italic, size, mono):
    txt = txt.replace(CURRENCY, "NT$")
    txt = txt.replace("\\|", "|").replace("\\*", "*").replace("\\_", "_")
    r = par.add_run(txt)
    r.bold = bold or None
    r.italic = italic or None
    if size:
        r.font.size = Pt(size)
    if mono:
        r.font.name = MONO_FONT
        r.font.size = Pt((size or BODY_PT) - 1.2)
    return r


# ------------------------------------------------------------------- block parsing


def protect(md: str) -> str:
    return md.replace("NT$", CURRENCY)


def split_cells(row: str) -> list[str]:
    row = row.strip().strip("|")
    out, cur, i = [], "", 0
    while i < len(row):
        if row[i] == "\\" and i + 1 < len(row):
            cur += row[i:i + 2]
            i += 2
            continue
        if row[i] == "|":
            out.append(cur.strip())
            cur = ""
            i += 1
            continue
        cur += row[i]
        i += 1
    out.append(cur.strip())
    return out


def build(md: str, doc: Document, *, hanging_refs=False) -> None:
    md = protect(md)
    lines = md.split("\n")
    i, n = 0, len(lines)
    eq_idx = 0
    in_appendix = False
    para: list[str] = []
    in_refs = False

    def flush():
        nonlocal para
        if not para:
            return
        # Strip the blockquote marker per line *before* joining: once the lines are
        # concatenated there are no line starts left for `^` to anchor to, and every
        # continuation marker would survive into the prose as a literal ">".
        quote = para[0].strip().startswith("> ")
        text = " ".join(re.sub(r"^>\s?", "", x.strip()) for x in para).strip()
        para = []
        if not text:
            return
        p = doc.add_paragraph()
        add_runs(p, text)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if quote:
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.15)
        elif in_refs or hanging_refs:
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)

    while i < n:
        line = lines[i]
        s = line.strip()

        # ---- HTML comment: a build marker, never content
        if s.startswith("<!--"):
            flush()
            while i < n and "-->" not in lines[i]:
                i += 1
            i += 1
            continue

        # ---- explicit page break
        if s.startswith('<div style="page-break-after'):
            flush()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            in_appendix = True
            i += 1
            continue

        # ---- fenced code
        if s.startswith("```"):
            flush()
            i += 1
            code: list[str] = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            for k, cl in enumerate(code):
                p = doc.add_paragraph()
                r = p.add_run(cl.replace("\t", "    ") or " ")
                r.font.name = MONO_FONT
                r.font.size = Pt(CODE_PT)
                pf = p.paragraph_format
                pf.space_before = Pt(3 if k == 0 else 0)
                pf.space_after = Pt(4 if k == len(code) - 1 else 0)
                pf.left_indent = Inches(0.12)
                pf.line_spacing = 0.92
                pf.keep_together = True
                shade(p._p.get_or_add_pPr(), "F4F4F4")
            continue

        # ---- display maths
        if s.startswith("$$"):
            flush()
            buf = [s]
            while not (len(buf) > 1 and buf[-1].rstrip().endswith("$$")) and \
                    not (buf[0].rstrip().endswith("$$") and len(buf[0]) > 4):
                i += 1
                if i >= n:
                    break
                buf.append(lines[i].strip())
            i += 1
            src = " ".join(buf).strip().strip("$")
            eq_idx += 1
            png = render_equation(src, eq_idx)
            w = min(TEXT_W - 0.6, 0.0075 * _px_width(png))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(png), width=Inches(w))
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(5)
            continue

        # ---- image
        if s.startswith("!["):
            flush()
            m = re.search(r"\(([^)]+)\)", s)
            if m:
                path = (ROOT / m.group(1)).resolve()
                if path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(
                        str(path),
                        width=Inches(FIG_W_APPX if in_appendix else FIG_W_BODY),
                    )
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = True
                else:
                    print(f"  !! missing image: {path}", file=sys.stderr)
            i += 1
            continue

        # ---- table
        if s.startswith("|"):
            flush()
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            grid = [split_cells(r) for r in rows]
            grid = [g for g in grid if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in g)]
            if not grid:
                continue
            ncol = max(len(g) for g in grid)
            t = doc.add_table(rows=0, cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = True
            for ri, g in enumerate(grid):
                cells = t.add_row().cells
                for ci in range(ncol):
                    txt = g[ci] if ci < len(g) else ""
                    cell = cells[ci]
                    cp = cell.paragraphs[0]
                    add_runs(cp, txt, bold=(ri == 0), size=TABLE_PT)
                    cp.paragraph_format.space_before = Pt(0.5)
                    cp.paragraph_format.space_after = Pt(0.5)
                    cp.paragraph_format.line_spacing = 0.95
                    if ri == 0:
                        shade(cell._tc.get_or_add_tcPr(), "EDEDED")
                    if ci > 1:
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            repeat_header(t.rows[0])
            tight_cell_margins(t)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # ---- headings
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush()
            lvl, txt = len(m.group(1)), m.group(2)
            in_refs = bool(re.match(r"^\d*\.?\s*References", txt)) or \
                (lvl == 2 and "References" in txt)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            if lvl == 1:
                add_runs(p, txt, bold=True, size=15)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf.space_before = Pt(0)
                pf.space_after = Pt(6)
            elif lvl == 2:
                add_runs(p, txt, bold=True, size=12)
                pf.space_before = Pt(9)
                pf.space_after = Pt(3)
            elif lvl == 3:
                add_runs(p, txt, bold=True, size=10.5)
                pf.space_before = Pt(7)
                pf.space_after = Pt(2)
            else:
                add_runs(p, txt, bold=True, italic=True, size=10)
                pf.space_before = Pt(6)
                pf.space_after = Pt(1)
            pf.keep_with_next = True
            i += 1
            continue

        # ---- horizontal rule / blank
        if s in ("---", "***", "___"):
            flush()
            i += 1
            continue
        if not s:
            flush()
            i += 1
            continue

        para.append(line)
        i += 1

    flush()


def _px_width(png: Path) -> int:
    from PIL import Image
    with Image.open(png) as im:
        return im.size[0]


# ----------------------------------------------------------------------- entry


def make(src: Path, dst: Path, *, hanging_refs=False) -> None:
    doc = Document()
    setup(doc)
    build(src.read_text(encoding="utf-8"), doc, hanging_refs=hanging_refs)
    doc.save(dst)
    print(f"wrote {dst.relative_to(REPO)}")


if __name__ == "__main__":
    make(ROOT / "phase2_report.md", ROOT / "phase2_report.docx")
    make(ROOT / "annotated_bibliography.md", ROOT / "annotated_bibliography.docx")
